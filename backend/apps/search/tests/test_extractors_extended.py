"""Estrattori XLSX, CSV, PPTX, EML (FASE 37)."""
import email
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook


def test_extract_xlsx(tmp_path):
    from apps.search.extractors import extract_text

    p = tmp_path / "t.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["A1", "B1"])
    ws.append([2, 3])
    wb.save(p)
    out = extract_text(str(p), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "A1" in out
    assert "B1" in out


def test_extract_csv(tmp_path):
    from apps.search.extractors import extract_text

    p = tmp_path / "d.csv"
    p.write_text("uno,due\ntre,quattro\n", encoding="utf-8")
    out = extract_text(str(p), "text/csv")
    assert "uno" in out
    assert "quattro" in out


def test_extract_pptx(tmp_path):
    from apps.search.extractors import extract_text

    p = tmp_path / "s.pptx"
    p.write_bytes(b"x")
    mock_slide = MagicMock()
    mock_shape = MagicMock()
    mock_shape.text = "slide text"
    mock_slide.shapes = [mock_shape]
    mock_prs = MagicMock()
    mock_prs.slides = [mock_slide]
    with patch("pptx.Presentation", return_value=mock_prs):
        out = extract_text(str(p), "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert "slide text" in out


def test_extract_eml(tmp_path):
    from apps.search.extractors import extract_text

    p = tmp_path / "m.eml"
    msg = email.message.EmailMessage()
    msg["Subject"] = "Oggetto test"
    msg["From"] = "a@b.it"
    msg.set_content("Corpo messaggio")
    p.write_bytes(msg.as_bytes())
    out = extract_text(str(p), "message/rfc822")
    assert "Oggetto test" in out
    assert "Corpo messaggio" in out


def test_extract_unknown_format_returns_empty(tmp_path):
    from apps.search.extractors import extract_text

    p = tmp_path / "z.bin"
    p.write_bytes(b"\x00\x01\x02")
    assert extract_text(str(p), "application/octet-stream") == ""
